"""Document loading and text extraction (agentic-research-intelligence-platform.md section 6)."""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

from app.core.logging import get_logger

logger = get_logger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown", ".csv", ".json", ".html", ".htm"}


@dataclass
class LoadedDocument:
    title: str
    source: str
    text: str
    category: str | None = None


def load_document(path: str | Path, category: str | None = None) -> LoadedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = _load_pdf(path)
    elif suffix == ".docx":
        text = _load_docx(path)
    elif suffix == ".csv":
        text = _load_csv(path)
    elif suffix == ".json":
        text = _load_json(path)
    elif suffix in {".html", ".htm"}:
        text = _load_html(path)
    elif suffix in {".txt", ".md", ".markdown"}:
        text = path.read_text(encoding="utf-8", errors="ignore")
    else:
        raise ValueError(f"Unsupported document type: {suffix} (supported: {SUPPORTED_EXTENSIONS})")

    return LoadedDocument(title=path.stem, source=str(path), text=_clean_text(text), category=category)


def load_documents_from_dir(directory: str | Path) -> list[LoadedDocument]:
    """Loads every supported file under `directory`, tagging each with its
    immediate subfolder name as `category` (e.g. `documents/financial/x.md`
    -> category "financial") - a free, zero-cost topic signal for datasets
    that are already organized this way.
    """
    directory = Path(directory)
    docs: list[LoadedDocument] = []
    for path in sorted(directory.rglob("*")):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            try:
                relative_parts = path.relative_to(directory).parts
                category = relative_parts[0] if len(relative_parts) > 1 else None
                docs.append(load_document(path, category=category))
            except Exception:  # noqa: BLE001
                logger.exception("Failed to load document %s", path)
    return docs


def _load_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _load_docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    return "\n\n".join(p.text for p in document.paragraphs)


def _load_csv(path: Path) -> str:
    with path.open("r", encoding="utf-8", errors="ignore", newline="") as f:
        rows = list(csv.reader(f))
    if not rows:
        return ""
    header, *body = rows
    return "\n".join(", ".join(f"{h}: {v}" for h, v in zip(header, row)) for row in body)


def _load_json(path: Path) -> str:
    data = json.loads(path.read_text(encoding="utf-8", errors="ignore"))
    return json.dumps(data, indent=2, ensure_ascii=False)


class _HTMLTextExtractor(HTMLParser):
    """Minimal tag stripper - avoids pulling in a full HTML parsing dependency."""

    def __init__(self) -> None:
        super().__init__()
        self._chunks: list[str] = []
        self._skip = False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self._skip = True

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self._skip = False

    def handle_data(self, data: str) -> None:
        if not self._skip and data.strip():
            self._chunks.append(data.strip())

    def get_text(self) -> str:
        return "\n".join(self._chunks)


def _load_html(path: Path) -> str:
    extractor = _HTMLTextExtractor()
    extractor.feed(path.read_text(encoding="utf-8", errors="ignore"))
    return extractor.get_text()


def _clean_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines()]
    lines = [line for line in lines if line]
    return "\n".join(lines)
