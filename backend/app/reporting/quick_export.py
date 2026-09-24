"""Instant .docx export of a finished research result (no LLM calls, no knowledge-base access).

Uses only standard python-docx features (built-in styles, plain tables, external hyperlinks) so the
file opens the same in Microsoft Word, LibreOffice, Google Docs and Pages. Because it never touches
retrieval, it cannot pause or otherwise disturb background document indexing.
"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from app.reporting.docx_builder import _add_hyperlink, _clean, _set_style_font

_BOLD_RE = re.compile(r"(\*\*[^*]+\*\*)")

SECTION_TITLES = [
    ("market_overview", "Overview"),
    ("competitor_analysis", "Competitor Analysis"),
]
LIST_SECTIONS = [
    ("key_findings", "Key Findings"),
    ("opportunities", "Opportunities"),
    ("risks", "Risks"),
]


def _rich(paragraph, text: str) -> None:
    for part in _BOLD_RE.split(_clean(text)):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def _paragraphs(doc, text: str) -> None:
    for block in re.split(r"\n\s*\n", _clean(text)):
        block = block.strip()
        if block:
            _rich(doc.add_paragraph(), block)


def _bullets(doc, items: list[Any], style: str = "List Bullet") -> None:
    for item in items:
        if str(item).strip():
            _rich(doc.add_paragraph(style=style), str(item))


def build_quick_docx(query: str, report: dict[str, Any], sources: list[dict[str, Any]], metrics: dict[str, Any] | None = None) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.2)
    section.top_margin = section.bottom_margin = Cm(2.0)

    _set_style_font(doc.styles["Normal"], "Calibri", 10.5)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    _set_style_font(doc.styles["Heading 1"], "Calibri", 16, True, "1F3864")
    _set_style_font(doc.styles["Heading 2"], "Calibri", 13, True, "2E75B6")

    title = doc.add_paragraph()
    run = title.add_run("Research Report")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    sub = doc.add_paragraph()
    sub_run = sub.add_run(_clean(query))
    sub_run.font.size = Pt(13)
    sub_run.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    stamp = doc.add_paragraph()
    stamp_run = stamp.add_run(f"Generated {datetime.now().astimezone():%d %B %Y, %H:%M}")
    stamp_run.italic = True
    stamp_run.font.size = Pt(9)

    if report.get("executive_summary"):
        doc.add_heading("Executive Summary", level=1)
        _paragraphs(doc, report["executive_summary"])

    for key, heading in SECTION_TITLES[:1]:
        if report.get(key):
            doc.add_heading(heading, level=1)
            _paragraphs(doc, report[key])

    for key, heading in LIST_SECTIONS[:1]:
        if report.get(key):
            doc.add_heading(heading, level=1)
            _bullets(doc, report[key])

    for key, heading in SECTION_TITLES[1:]:
        if report.get(key):
            doc.add_heading(heading, level=1)
            _paragraphs(doc, report[key])

    for key, heading in LIST_SECTIONS[1:]:
        if report.get(key):
            doc.add_heading(heading, level=1)
            _bullets(doc, report[key])

    if report.get("evidence"):
        doc.add_heading("Supporting Evidence", level=1)
        _bullets(doc, report["evidence"])

    verification = report.get("verification") or {}
    notes = list(verification.get("low_confidence_claims") or [])
    if notes:
        doc.add_heading("Points to Verify", level=1)
        doc.add_paragraph("These claims are supported only weakly (single source or dated figure):")
        _bullets(doc, notes)

    if sources:
        doc.add_heading("Sources", level=1)
        for i, src in enumerate(sources, start=1):
            para = doc.add_paragraph(style="List Number")
            para.add_run(_clean(src.get("title") or "Untitled"))
            origin = {"web": "web", "knowledge_base": "knowledge base", "uploaded_document": "uploaded document"}.get(
                src.get("origin", ""), ""
            )
            if origin:
                tail = para.add_run(f"  ({origin})")
                tail.italic = True
                tail.font.size = Pt(9)
            url = src.get("url")
            if url and str(url).startswith(("http://", "https://")):
                para.add_run("  ")
                _add_hyperlink(para, str(url), str(url), size=9)

    if metrics:
        note = doc.add_paragraph()
        note.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = note.add_run(
            f"Run details: {metrics.get('llm_calls', 0)} model calls, "
            f"{metrics.get('tool_calls', 0)} tool calls, {round((metrics.get('latency_ms') or 0) / 1000)} s."
        )
        run.italic = True
        run.font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
